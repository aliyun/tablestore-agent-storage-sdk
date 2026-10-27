# Optional regeneration tool; not used by Go tests.
# Requires python-docx, openpyxl, python-pptx, reportlab and Pillow.
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas
from PIL import Image
p=Path(__file__).resolve().parent;p.mkdir(parents=True,exist_ok=True)
d=Document();d.core_properties.author='SDK acceptance';d.add_paragraph('Synthetic SDK fixture. Project Cedar uses code 4827.');d.save(p/'source.docx')
w=Workbook();w.active.append(['Project','Code']);w.active.append(['Cedar',4827]);w.properties.creator='SDK acceptance';w.save(p/'source.xlsx')
r=Presentation();s=r.slides.add_slide(r.slide_layouts[6]);s.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(2)).text='Synthetic SDK fixture: Cedar code 4827';r.core_properties.author='SDK acceptance';r.save(p/'source.pptx')
c=Canvas(str(p/'source.pdf'),invariant=1);c.setAuthor('SDK acceptance');c.drawString(72,720,'Synthetic SDK fixture: Cedar code 4827');c.save()
for ext in ['png','jpg','jpeg']: Image.new('RGB',(8,8),(12,34,56)).save(p/('source.'+ext))
